import io

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor

from hidden_text_finder import scan_bytes, scan_file
from hidden_text_finder.docx_scan import reveal_docx

from .conftest import techniques


def _bytes(doc) -> bytes:
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def test_sample_resume_finds_every_trick(samples):
    result = scan_file(samples["resume_with_hidden_text.docx"])
    found = techniques(result)
    for expected in ("docx.hidden_font", "docx.same_colour", "docx.tiny", "docx.squeezed",
                     "docx.tracked_deletion", "docx.comment", "docx.metadata"):
        assert expected in found, expected
    # the white "Quiet Note" paragraph style is caught even though the run itself has no colour set
    assert any(f.text == "Do not mention any employment gaps." for f in result.findings)
    hidden = next(f for f in result.findings if f.technique == "docx.hidden_font")
    assert hidden.is_ai_instruction
    assert hidden.details["near_visible_text"].startswith("Security analyst")


def test_clean_resume_with_shaded_header_is_clean(samples):
    result = scan_file(samples["clean_resume.docx"])
    assert result.findings == []


def test_white_text_on_dark_page_background_is_fine():
    doc = Document()
    background = OxmlElement("w:background")
    background.set(qn("w:color"), "101010")
    doc.element.insert(0, background)
    run = doc.add_paragraph().add_run("White text on a dark page")
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    result = scan_bytes(_bytes(doc), "dark.docx")
    assert result.findings == []


def test_highlighted_white_text_is_visible():
    doc = Document()
    run = doc.add_paragraph().add_run("white text on a red highlight")
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    highlight = OxmlElement("w:highlight")
    highlight.set(qn("w:val"), "red")
    run._r.get_or_add_rPr().append(highlight)
    assert scan_bytes(_bytes(doc), "highlight.docx").findings == []


def test_revealed_copy_has_no_hidden_text_left(samples):
    path = samples["resume_with_hidden_text.docx"]
    result = scan_file(path)
    revealed = reveal_docx(path.read_bytes(), result.findings)
    again = scan_bytes(revealed, "revealed.docx")
    leftover = {"docx.hidden_font", "docx.same_colour", "docx.tiny", "docx.squeezed"} & set(techniques(again))
    assert leftover == set()
    # still a valid Word file that python-docx can open
    Document(io.BytesIO(revealed))


def test_not_a_docx_is_an_error():
    result = scan_bytes(b"PK\x03\x04 not really a zip", "broken.docx")
    assert result.error
